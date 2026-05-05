from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import openpyxl

from app.infrastructure.db.repositories.measurement_repo import MeasurementRepository
from app.infrastructure.db.repositories.part_repo import PartRepository
from app.infrastructure.db.repositories.template_parameter_repo import TemplateParameterRepository


def _safe_sheet_name(name: str) -> str:
    """Return a valid Excel sheet name (max 31 chars, no forbidden chars)."""
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name)
    return cleaned[:31]


def _safe_filename(name: str) -> str:
    """Return a filename-safe string."""
    return re.sub(r'[<>:"/\\|?*]', "_", name)


class ArchiveService:
    def __init__(self) -> None:
        self._part_repo = PartRepository()
        self._param_repo = TemplateParameterRepository()
        self._meas_repo = MeasurementRepository()

    def export(self, directory: str, archive_name: str) -> Path:
        """
        Build a ZIP archive at ``directory / archive_name.zip``.

        The archive contains one .xlsx file per part.  Each sheet inside
        an xlsx file corresponds to one template parameter and contains all
        measurements for that part + parameter (columns: operating_hours, value).

        Returns the full path to the created archive.
        """
        dest_dir = Path(directory)
        dest_dir.mkdir(parents=True, exist_ok=True)

        stem = archive_name if archive_name.lower().endswith(".zip") else f"{archive_name}.zip"
        archive_path = dest_dir / stem

        parts = list(self._part_repo.list())

        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for part in parts:
                xlsx_bytes = self._build_part_xlsx(part.id, part.template_id)
                if part.serial_number:
                    filename = _safe_filename(f"{part.name}_{part.serial_number}") + ".xlsx"
                else:
                    filename = _safe_filename(part.name) + ".xlsx"
                zf.writestr(filename, xlsx_bytes)

        return archive_path

    def _build_part_xlsx(self, part_id: int, template_id: int) -> bytes:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default empty sheet

        params = list(self._param_repo.list_by_template(template_id))

        for param in params:
            measurements = list(self._meas_repo.list_by_part_and_parameter(part_id, param.id))
            sheet_name = _safe_sheet_name(param.name)
            ws = wb.create_sheet(title=sheet_name)

            ws.append(["Часы наработки", f"Значение ({param.unit})"])
            for m in measurements:
                ws.append([m.operating_hours, m.value])

        if not wb.sheetnames:
            ws = wb.create_sheet(title="Нет данных")
            ws.append(["Нет параметров для данной детали"])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
